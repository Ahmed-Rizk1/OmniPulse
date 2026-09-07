import io
from typing import Sequence
import uuid
from fastapi import HTTPException, UploadFile, status
from pypdf import PdfReader
import docx
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession
import structlog

from app.modules.knowledge.models import DocumentChunk, SourceDocument
from app.modules.knowledge.schemas import DocumentUploadOut, KnowledgeDocumentOut
from app.shared.ai import get_embedding_client
from app.shared.database.rls import set_tenant_context

logger = structlog.get_logger("app.modules.knowledge.service")


def extract_text_from_file(filename: str, content: bytes) -> str:
    """Extracts raw text from PDF, DOCX, TXT, or MD files."""
    lower = filename.lower()
    if lower.endswith(".pdf"):
        try:
            reader = PdfReader(io.BytesIO(content))
            pages_text = [page.extract_text() or "" for page in reader.pages]
            return "\n\n".join(pages_text).strip()
        except Exception as exc:
            logger.error("pdf_extraction_error", filename=filename, error=str(exc))
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Failed to parse PDF file: {exc}",
            )
    elif lower.endswith(".docx"):
        try:
            doc = docx.Document(io.BytesIO(content))
            paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            return "\n\n".join(paras).strip()
        except Exception as exc:
            logger.error("docx_extraction_error", filename=filename, error=str(exc))
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Failed to parse DOCX file: {exc}",
            )
    elif lower.endswith(".txt") or lower.endswith(".md"):
        try:
            return content.decode("utf-8", errors="ignore").strip()
        except Exception as exc:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Failed to read text file: {exc}",
            )
    else:
        # Fallback to UTF-8 decoding
        try:
            return content.decode("utf-8", errors="ignore").strip()
        except Exception:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Unsupported file format. Allowed: .pdf, .docx, .txt, .md",
            )


def chunk_text(text: str, target_chunk_size: int = 700, overlap: int = 100) -> list[str]:
    """Chunks text into readable segments respecting paragraph boundaries."""
    if not text:
        return []

    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    if not paragraphs:
        paragraphs = [text.strip()]

    chunks: list[str] = []
    current_chunk: list[str] = []
    current_len = 0

    for para in paragraphs:
        para_len = len(para)
        # If single paragraph is longer than target, break it down
        if para_len > target_chunk_size:
            if current_chunk:
                chunks.append("\n\n".join(current_chunk))
                current_chunk = []
                current_len = 0

            start = 0
            while start < para_len:
                end = min(start + target_chunk_size, para_len)
                chunks.append(para[start:end])
                start += target_chunk_size - overlap
            continue

        if current_len + para_len > target_chunk_size:
            chunks.append("\n\n".join(current_chunk))
            current_chunk = [para]
            current_len = para_len
        else:
            current_chunk.append(para)
            current_len += para_len + 2

    if current_chunk:
        chunks.append("\n\n".join(current_chunk))

    return [c.strip() for c in chunks if c.strip()]


async def list_documents(
    tenant_id: uuid.UUID,
    db: AsyncSession,
) -> list[KnowledgeDocumentOut]:
    """Retrieves all source documents for the tenant."""
    await set_tenant_context(db, tenant_id)
    stmt = (
        select(SourceDocument)
        .where(SourceDocument.tenant_id == tenant_id)
        .order_by(SourceDocument.created_at.desc())
    )
    result = await db.execute(stmt)
    docs = result.scalars().all()

    return [
        KnowledgeDocumentOut(
            doc_id=doc.id,
            title=doc.title,
            file_type=doc.file_type.upper(),
            chunk_count=doc.chunk_count,
            created_at=doc.created_at,
        )
        for doc in docs
    ]


async def upload_and_process_document(
    file: UploadFile,
    tenant_id: uuid.UUID,
    db: AsyncSession,
) -> DocumentUploadOut:
    """Extracts text from file, chunks it, generates Gemini embeddings, and persists."""
    filename = file.filename or "untitled.txt"
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else "txt"
    if ext not in ("pdf", "docx", "txt", "md"):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"File extension '.{ext}' is not supported. Please upload .pdf, .docx, .txt, or .md",
        )

    content = await file.read()
    if not content:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Uploaded file is empty",
        )

    raw_text = extract_text_from_file(filename, content)
    if not raw_text:
        raw_text = f"Empty document: {filename}"

    chunks = chunk_text(raw_text)
    if not chunks:
        chunks = [raw_text[:1000]]

    # Generate embeddings for each chunk
    embedding_client = get_embedding_client()
    embeddings: list[list[float] | None] = []

    try:
        batch_results = await embedding_client.embed_batch(chunks)
        embeddings = batch_results  # type: ignore
    except Exception as exc:
        logger.warn("gemini_embed_batch_fallback_individual", error=str(exc))
        for chunk in chunks:
            try:
                emb = await embedding_client.embed_text(chunk)
                embeddings.append(emb)
            except Exception as single_exc:
                logger.error("chunk_embedding_failed", error=str(single_exc))
                embeddings.append(None)

    # Persist document and chunks
    await set_tenant_context(db, tenant_id)
    doc = SourceDocument(
        id=uuid.uuid4(),
        tenant_id=tenant_id,
        title=filename,
        file_type=ext.upper(),
        chunk_count=len(chunks),
    )
    db.add(doc)
    await db.flush()

    for idx, (chunk_text_content, emb) in enumerate(zip(chunks, embeddings)):
        chunk_row = DocumentChunk(
            id=uuid.uuid4(),
            tenant_id=tenant_id,
            document_id=doc.id,
            chunk_index=idx,
            content=chunk_text_content,
            embedding=emb,
        )
        db.add(chunk_row)

    await db.commit()
    logger.info(
        "document_processed_and_saved",
        doc_id=str(doc.id),
        tenant_id=str(tenant_id),
        title=doc.title,
        chunk_count=doc.chunk_count,
    )

    return DocumentUploadOut(
        doc_id=doc.id,
        title=doc.title,
        chunk_count=doc.chunk_count,
    )


async def delete_document(
    doc_id: uuid.UUID,
    tenant_id: uuid.UUID,
    db: AsyncSession,
) -> bool:
    """Deletes a source document and cascades to all its chunks."""
    await set_tenant_context(db, tenant_id)
    stmt = (
        delete(SourceDocument)
        .where(
            SourceDocument.id == doc_id,
            SourceDocument.tenant_id == tenant_id,
        )
    )
    result = await db.execute(stmt)
    await db.commit()

    deleted = (result.rowcount or 0) > 0
    if not deleted:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Document {doc_id} not found",
        )

    logger.info(
        "document_deleted",
        doc_id=str(doc_id),
        tenant_id=str(tenant_id),
    )
    return True
